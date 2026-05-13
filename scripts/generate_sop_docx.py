#!/usr/bin/env python3
"""
Builds Beach-Aware-SOP-and-GitHub-Guide.docx from structured content.
Run: python3 scripts/generate_sop_docx.py
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def shade_cell(cell, fill_hex: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill_hex)
    cell._tc.get_or_add_tcPr().append(shd)


def heading(doc: Document, text: str, level: int) -> None:
    doc.add_heading(text, level=level)


def para(doc: Document, text: str, *, bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(11)


def bullet(doc: Document, text: str, *, level: int = 0) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 + level * 0.25)
    for run in p.runs:
        run.font.size = Pt(11)


def numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="List Number")
    for run in p.runs:
        run.font.size = Pt(11)


def code_block(doc: Document, text: str) -> None:
    for line in text.strip().split("\n"):
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_together = True


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    *,
    header_fill: str = "D9E2F3",
) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        shade_cell(hdr[i], header_fill)
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
            for p in cells[ci].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()


def horizontal_rule(doc: Document) -> None:
    doc.add_paragraph()


def build() -> Document:
    doc = Document()
    core = doc.core_properties
    core.title = "Beach Aware — SOP & GitHub Guide"
    core.subject = "Devstars Jersey — Standard operating procedures and version control"
    core.comments = "Generated for Beach Aware project contributors."

    # --- Title block ---
    t = doc.add_heading("Beach Aware Project", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("Standard Operating Procedures & GitHub User Guide")
    r.bold = True
    r.font.size = Pt(14)
    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run("Devstars Jersey\nDocument version 1.0 • May 2026")
    mr.font.size = Pt(10)
    mr.italic = True
    doc.add_page_break()

    # --- Document control ---
    heading(doc, "Document control", 1)
    para(
        doc,
        "This document explains how the Beach Aware project is organised, how to work on it safely, "
        "and how to use GitHub together with the codebase. It is written for developers, "
        "technical staff, and anyone who needs to update or deploy the application.",
    )
    add_table(
        doc,
        ["Field", "Detail"],
        [
            ["Project name", "Beach Aware (Devstars Jersey)"],
            ["Purpose", "SOP for development workflow; GitHub usage; linkage to app architecture"],
            ["Primary repository", "git@github.com:emmanuel-devstar/beach-aware-project.git"],
            ["Default branch", "main"],
            ["Companion docs", "README.md (repository root), php-tuner/README.md (PHP tuner setup)"],
        ],
    )
    horizontal_rule(doc)

    heading(doc, "Table of contents", 1)
    toc_items = [
        "1. Executive summary",
        "2. Project overview and architecture",
        "3. Repository layout",
        "4. Standard operating procedures (daily work)",
        "5. Release and deployment checklist (high level)",
        "6. GitHub: concepts and setup",
        "7. GitHub: day-to-day commands",
        "8. Branching and pull requests (recommended practice)",
        "9. Merge conflicts",
        "10. Security, secrets, and .gitignore",
        "11. Troubleshooting",
        "12. Command reference (cheat sheet)",
    ]
    for item in toc_items:
        bullet(doc, item)
    doc.add_page_break()

    # --- 1 Executive summary ---
    heading(doc, "1. Executive summary", 1)
    para(
        doc,
        "Beach Aware is a beta web application for Jersey that helps users understand beach and sea "
        "conditions for activities such as swimming, paddling, and other water sports. "
        "The public site is a static front-end. A separate PHP application (the “tuner”) stores "
        "adjustments in a database and exposes read-only JSON via api.php for the front-end.",
    )
    para(
        doc,
        "GitHub holds the shared source code. Secrets (database passwords, session keys) must never "
        "be committed. Deployments should use the latest agreed code on the main branch unless your "
        "team uses another process.",
    )

    # --- 2 Architecture ---
    heading(doc, "2. Project overview and architecture", 1)
    heading(doc, "2.1 How the pieces fit together", 2)
    numbered(doc, "A visitor opens index.html in a browser. The page loads JavaScript and CSS from the assets/ folder.")
    numbered(doc, "On startup, the application requests JSON from the PHP tuner API (api.php) using the URL configured in index.html (window.__BEACH_AWARE_TUNER_API__).")
    numbered(doc, "Authorised staff use the tuner’s admin pages (after login) to edit per-beach, per-activity settings stored in MySQL/MariaDB.")
    numbered(doc, "If the API is unavailable, the front-end falls back to default values embedded in the build so users still see a working page.")
    heading(doc, "2.2 Roles (informal)", 2)
    add_table(
        doc,
        ["Role", "Typical tasks"],
        [
            ["Developer", "Change code, run Git commands, open pull requests, coordinate deploys"],
            ["Content / ops (tuner)", "Log in to PHP admin; adjust biases and notes; avoid raw SQL unless trained"],
            ["Lead / reviewer", "Approve merges, resolve escalations, own production config"],
        ],
    )

    # --- 3 Repo layout ---
    heading(doc, "3. Repository layout", 1)
    para(doc, "Key paths in the beach-aware-project repository:", bold=True)
    add_table(
        doc,
        ["Path", "Description"],
        [
            ["index.html", "Entry HTML; sets API URL for the tuner; loads the built React bundle."],
            ["assets/", "Compiled JavaScript, CSS, images for the public site."],
            ["php-tuner/", "PHP application: public web root (api.php, login, admin), src/, sql/ schema."],
            ["php-tuner/.env", "Local secrets — present only on each machine/server; not in Git."],
            ["SOP-AND-GITHUB-GUIDE.md", "Markdown copy of this guidance (editable in Git)."],
            ["README.md", "Short project intro and links."],
        ],
    )

    # --- 4 SOP daily ---
    heading(doc, "4. Standard operating procedures (daily work)", 1)
    heading(doc, "4.1 Before you edit files", 2)
    numbered(doc, "Open a terminal and change to your local clone of the repository.")
    numbered(doc, "Fetch and integrate the latest main branch:")
    code_block(doc, "git checkout main\ngit pull origin main")
    numbered(doc, "Confirm you are on the correct task (ticket, message from lead, etc.).")

    heading(doc, "4.2 While you work", 2)
    bullet(doc, "Save files frequently.")
    bullet(doc, "Do not copy real .env values into tickets, email, or screenshots.")
    bullet(doc, "If you add new secret files, add their names to .gitignore and tell the team how to create them locally.")

    heading(doc, "4.3 After a coherent change (commit and push)", 2)
    numbered(doc, "Review changes: git status and optionally git diff.")
    numbered(doc, "Stage files (example: all changes in the repo):")
    code_block(doc, "git add .")
    numbered(doc, "Commit with a clear message:")
    code_block(doc, 'git commit -m "Describe the change in one short line"')
    numbered(doc, "Upload to GitHub:")
    code_block(doc, "git push origin main")
    para(doc, "If your team uses feature branches, push your branch instead and open a pull request (section 8).", italic=True)

    heading(doc, "4.4 End of day", 2)
    bullet(doc, "Push or hand off incomplete work on a branch; do not leave important changes only on one laptop.")
    bullet(doc, "Note blockers in your team chat or ticket system.")

    # --- 5 Deploy checklist ---
    heading(doc, "5. Release and deployment checklist (high level)", 1)
    para(
        doc,
        "Exact steps depend on your hosting. Use this as a common checklist; adapt with your "
        "infrastructure runbook.",
    )
    numbered(doc, "Confirm main (or release branch) contains the tested commit.")
    numbered(doc, "Build or copy the static front-end (index.html, assets/) if your process rebuilds from React source.")
    numbered(doc, "Deploy php-tuner files and ensure .env on the server has correct DB_* and security values (never from Git).")
    numbered(doc, "Confirm index.html on the server points __BEACH_AWARE_TUNER_API__ to the live api.php URL.")
    numbered(doc, "Smoke-test: public page loads; tuner login works; API returns JSON; a test edit appears in the app after cache period.")

    # --- 6 GitHub concepts ---
    heading(doc, "6. GitHub: concepts and setup", 1)
    heading(doc, "6.1 Key terms", 2)
    add_table(
        doc,
        ["Term", "Meaning"],
        [
            ["Git", "Version control software; runs on your computer."],
            ["GitHub", "Cloud service that stores the Git repository and enables collaboration."],
            ["Repository (repo)", "The project folder tracked by Git; hosted at the URL above."],
            ["Clone", "One-time copy of the repo from GitHub to your machine."],
            ["Commit", "A saved snapshot of changes on your computer, with a message."],
            ["Push", "Upload your new commits to GitHub."],
            ["Pull", "Download and merge others’ commits from GitHub into your branch."],
            ["Branch", "A separate line of development (e.g. main, feature/xyz)."],
            ["Remote", "Named link to GitHub; origin usually points to the primary repo."],
        ],
    )
    heading(doc, "6.2 SSH access (typical Devstars / GitHub setup)", 2)
    para(
        doc,
        "This project uses an SSH remote URL (git@github.com:…). Your Mac should have an SSH key "
        "registered with GitHub and a Host entry in ~/.ssh/config if your organisation uses a "
        "specific identity file. If clone or push asks for a password or fails with “Permission denied”, "
        "check SSH agent and key configuration with your IT lead.",
    )
    code_block(doc, "ssh -T git@github.com\n# Expect a greeting mentioning your GitHub username if SSH works.")

    heading(doc, "6.3 One-time clone", 2)
    code_block(
        doc,
        "cd ~/Documents    # or your preferred folder\ngit clone git@github.com:emmanuel-devstar/beach-aware-project.git\ncd beach-aware-project",
    )

    # --- 7 Day-to-day Git ---
    heading(doc, "7. GitHub: day-to-day commands", 1)
    para(doc, "Typical session:", bold=True)
    code_block(
        doc,
        "cd /path/to/beach-aware-project\ngit checkout main\ngit pull origin main\n# edit files\ngit status\ngit add .\ngit commit -m \"Your message\"\ngit push origin main",
    )
    para(
        doc,
        "Use git diff before committing to catch accidental edits (especially to .env if "
        "it were ever tracked).",
    )

    # --- 8 Branching & PRs ---
    heading(doc, "8. Branching and pull requests (recommended practice)", 1)
    para(
        doc,
        "For anything non-trivial, create a branch from main, push it, and open a Pull Request (PR) "
        "on GitHub. A reviewer can comment before code reaches main.",
    )
    code_block(
        doc,
        "git checkout main\ngit pull origin main\ngit checkout -b feature/short-description\n# work, commit locally\ngit push -u origin feature/short-description",
    )
    bullet(doc, "On github.com, open a Pull Request from feature/short-description into main.")
    bullet(doc, "After approval, use “Merge” on GitHub (or merge locally if that is team policy).")
    bullet(doc, "Delete the remote branch after merge if no longer needed.")

    # --- 9 Conflicts ---
    heading(doc, "9. Merge conflicts", 1)
    para(
        doc,
        "A conflict happens when two people change the same lines. Git cannot merge automatically.",
    )
    numbered(doc, "Run: git pull origin main (on your branch) or complete the merge/ rebase your team uses.")
    numbered(doc, "Open files listed as conflicted. Search for markers: <<<<<<<, =======, >>>>>>>.")
    numbered(doc, "Edit to the final desired text; remove all marker lines.")
    numbered(doc, "Stage resolved files: git add <file>")
    numbered(doc, "Complete the merge: git commit (if Git opened a merge) or continue rebase per instructions.")
    numbered(doc, "Push the result. Avoid git push --force to main unless explicitly authorised.")
    para(doc, "Escalate to a colleague the first time you resolve conflicts — it is a normal learning step.", italic=True)

    # --- 10 Security ---
    heading(doc, "10. Security, secrets, and .gitignore", 1)
    bullet(doc, "php-tuner/.env holds database credentials, session secret, and admin password hash. It must remain local or on the server only.")
    bullet(doc, ".gitignore is configured to exclude .env; verify with git status before each commit.")
    bullet(doc, "If secrets were ever pushed to GitHub, rotate all exposed values immediately and contact your lead.")
    bullet(doc, "api.php is intentionally public read-only; the admin interface requires login — keep tuner URLs and passwords confidential.")

    # --- 11 Troubleshooting ---
    heading(doc, "11. Troubleshooting", 1)
    add_table(
        doc,
        ["Symptom", "What to try first"],
        [
            ["git push rejected (non-fast-forward)", "git pull origin main, resolve conflicts, push again."],
            ["Permission denied (publickey)", "SSH key not loaded or wrong GitHub account; check ssh -T git@github.com"],
            ["Public site ignores tuner changes", "Check API URL in index.html; clear browser cache; wait for front-end cache TTL (e.g. ~5 min per app design)."],
            ["Tuner login fails", "Verify .env on server; database reachable; bcrypt hash correct."],
            ["Blank or broken page", "Browser console for JS errors; verify assets paths and that index.html references correct script names."],
        ],
    )
    para(doc, "For database and Nginx detail, see php-tuner/README.md in the repository.", italic=True)

    # --- 12 Cheat sheet ---
    heading(doc, "12. Command reference (cheat sheet)", 1)
    code_block(
        doc,
        """git status                  # Files changed
git diff                    # Line-by-line changes
git log --oneline -10      # Recent commits
git checkout main          # Switch to main branch
git pull origin main       # Update main from GitHub
git add .                  # Stage all changes in folder
git commit -m \"msg\"       # Save commit locally
git push origin main       # Upload commits
git branch                 # List branches
git checkout -b name     # New branch from current HEAD""",
    )

    heading(doc, "Related files in the repository", 1)
    bullet(doc, "README.md — project summary")
    bullet(doc, "SOP-AND-GITHUB-GUIDE.md — Markdown version of this guide")
    bullet(doc, "php-tuner/README.md — MySQL schema, .env, Nginx, security")

    doc.add_paragraph()
    footer = doc.add_paragraph()
    fr = footer.add_run(
        "End of document. In Microsoft Word you can insert an automatic Table of Contents "
        "(References → Table of Contents) from the heading styles if you need page numbers updated."
    )
    fr.italic = True
    fr.font.size = Pt(9)

    return doc


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    out = root / "Beach-Aware-SOP-and-GitHub-Guide.docx"
    doc = build()
    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
